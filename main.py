from __future__ import annotations

import argparse
import base64
import csv
import random
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import arabic_reshaper
from bidi.algorithm import get_display
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from playwright.sync_api import sync_playwright

import config


def shape_for_image(text: str) -> str:
    """Shape Arabic letters and apply bidi ordering for raster image rendering."""
    return get_display(arabic_reshaper.reshape(text))


def load_recipients():
    if not config.EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Excel file not found: {config.EXCEL_FILE}\n"
            "Create input/recipients.xlsx with headers: name, phone."
        )

    workbook = load_workbook(config.EXCEL_FILE, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        raise ValueError("The Excel worksheet is empty.")

    normalized = [str(h).strip().lower() if h is not None else "" for h in headers]
    if "name" not in normalized or "phone" not in normalized:
        raise ValueError("Excel must contain 'name' and 'phone' column headers.")

    name_i = normalized.index("name")
    phone_i = normalized.index("phone")
    recipients = []
    for excel_row, row in enumerate(rows, start=2):
        name = str(row[name_i]).strip() if name_i < len(row) and row[name_i] is not None else ""
        raw_phone = row[phone_i] if phone_i < len(row) else None
        phone = str(raw_phone).strip() if raw_phone is not None else ""
        # Avoid converting numeric phone values like 9.05e+11 into malformed numbers.
        if isinstance(raw_phone, (int, float)):
            phone = str(int(raw_phone)) if float(raw_phone).is_integer() else str(raw_phone)

        phone = re.sub(r"[^\d+]", "", phone)
        if phone.startswith("+"):
            phone = phone[1:]
        if not name and not phone:
            continue
        if not name or not phone:
            print(f"Skipping Excel row {excel_row}: name or phone is empty.")
            continue
        if not phone.isdigit():
            print(f"Skipping Excel row {excel_row}: phone must contain country code and digits only.")
            continue
        message = config.DEFAULT_MESSAGE.replace("{name}", name)
        recipients.append({"name": name, "phone": phone, "message": message, "excel_row": excel_row})

    workbook.close()
    return recipients


def render_invitation(name: str, destination: Path):
    if not config.INVITATION_IMAGE.exists():
        raise FileNotFoundError(f"Invitation image not found: {config.INVITATION_IMAGE}")

    if not config.FONT_PATH.exists():
        raise FileNotFoundError(f"Font file not found: {config.FONT_PATH}")

    image = Image.open(config.INVITATION_IMAGE).convert("RGBA")
    draw = ImageDraw.Draw(image)

    display_name = name
    max_width = int(image.width * config.MAX_NAME_WIDTH_RATIO)

    def make_font(size):
        return ImageFont.truetype(
        str(config.FONT_PATH), size, layout_engine=ImageFont.Layout.RAQM
    )

    lo, hi = config.MIN_FONT_SIZE, config.FONT_SIZE
    font_size = lo
    while lo <= hi:
        mid = (lo + hi) // 2
        if draw.textlength(display_name, font=make_font(mid), language="ar") <= max_width:
            font_size = mid
            lo = mid + 1
        else:
            hi = mid - 1

    font = make_font(font_size)

    center_x = image.width // 2 + config.NAME_OFFSET_X
    center_y = image.height // 2 + config.NAME_OFFSET_Y

    print(f"Image size: {image.width} x {image.height}")
    print(f"Drawing name: {name!r} at ({center_x}, {center_y}), font size {font_size}")

    draw.text(
        (center_x, center_y),
        display_name,
        font=font,
        fill=config.NAME_COLOR,
        anchor=config.NAME_ANCHOR,
        language="ar",
    )

    destination.parent.mkdir(parents=True, exist_ok=True)
    image.convert("RGB").save(destination, format="PNG")


def personalized_path(name: str, index: int) -> Path:
    safe_name = re.sub(r"[^\w.-]+", "_", name, flags=re.UNICODE).strip("._")
    if not safe_name:
        safe_name = "recipient"
    return config.PERSONALIZED_DIR / f"{index:04d}_{safe_name}.png"


def command_preview(name: str):
    render_invitation(name, config.PREVIEW_FILE)
    print(f"Preview saved to: {config.PREVIEW_FILE}")


def command_prepare(recipients):
    config.PERSONALIZED_DIR.mkdir(parents=True, exist_ok=True)
    for index, recipient in enumerate(recipients, start=1):
        path = personalized_path(recipient["name"], index)
        render_invitation(recipient["name"], path)
        print(f"[{index}/{len(recipients)}] Prepared {recipient['name']}: {path}")
    print(f"\nPrepared {len(recipients)} image(s). Original invitation was not changed.")


def append_log(recipient, status, detail=""):
    config.LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    exists = config.LOG_FILE.exists()
    with config.LOG_FILE.open("a", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=["timestamp", "name", "phone", "status", "detail"])
        if not exists:
            writer.writeheader()
        writer.writerow({
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "name": recipient["name"],
            "phone": recipient["phone"],
            "status": status,
            "detail": detail,
        })


def send_one(page, recipient):
    """Open the chat, paste the image, type the caption, press Enter."""
    url = f"https://web.whatsapp.com/send?phone={recipient['phone']}"
    page.goto(url, wait_until="domcontentloaded", timeout=config.PAGE_TIMEOUT_MS)

    composer = page.locator('footer [contenteditable="true"]').first
    composer.wait_for(state="visible", timeout=config.PAGE_TIMEOUT_MS)
    time.sleep(1.5)

    # Paste the image so WhatsApp opens the normal photo preview.
    image_b64 = base64.b64encode(Path(recipient["image_path"]).read_bytes()).decode()
    composer.click()
    composer.evaluate(
        """(el, b64) => {
            const bytes = Uint8Array.from(atob(b64), c => c.charCodeAt(0));
            const file = new File([bytes], "invitation.png", { type: "image/png" });
            const dt = new DataTransfer();
            dt.items.add(file);
            el.dispatchEvent(new ClipboardEvent("paste", {
                clipboardData: dt, bubbles: true, cancelable: true
            }));
        }""",
        image_b64,
    )
    time.sleep(2.5)

    # Type the caption into the preview's caption box and send with Enter.
    caption_box = page.locator('div[contenteditable="true"]:not(footer *)').last
    caption_box.wait_for(state="visible", timeout=15000)
    caption_box.click()
    page.keyboard.insert_text(recipient["message"])
    time.sleep(0.5)
    caption_box.press("Enter")

    # Confirm the preview closed; otherwise treat it as a failure.
    caption_box.wait_for(state="detached", timeout=15000)
    time.sleep(2.0)


def command_send(recipients):
    if not recipients:
        print("No valid recipients found in Excel.")
        return

    # Prepare images first, before opening WhatsApp.
    config.PERSONALIZED_DIR.mkdir(parents=True, exist_ok=True)
    for index, recipient in enumerate(recipients, start=1):
        path = personalized_path(recipient["name"], index)
        render_invitation(recipient["name"], path)
        recipient["image_path"] = path
        
    print("WhatsApp Web will open. Log in by scanning the QR code if requested.")
    print("Messages are sent automatically. Press Ctrl+C at any time to stop.")
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=False)
        context = browser.new_context()
        page = context.new_page()
        page.goto(
            "https://web.whatsapp.com/",
            wait_until="commit",
            timeout=config.PAGE_TIMEOUT_MS,
        )
        input("When WhatsApp Web is fully loaded and logged in, press Enter here...")

        try:
            for index, recipient in enumerate(recipients, start=1):
                print(f"\n[{index}/{len(recipients)}] {recipient['name']} — +{recipient['phone']}")

                try:
                    send_one(page, recipient)
                    append_log(recipient, "sent", "Image with caption sent.")
                    print("Sent.")
                except Exception as error:
                    append_log(recipient, "error", str(error)[:500])
                    print(f"Could not complete this recipient: {error}")

                # Pause between recipients (skip after the last one).
                if index < len(recipients):
                    delay = random.uniform(config.MIN_DELAY_SECONDS, config.MAX_DELAY_SECONDS)
                    print(f"Waiting {delay:.0f}s before the next recipient...")
                    time.sleep(delay)
        except KeyboardInterrupt:
            print("\nStopped by user.")
        finally:
            print(f"Log file: {config.LOG_FILE}")
            input("Press Enter to close the browser...")
            context.close()
            browser.close()


def main():
    parser = argparse.ArgumentParser(description="Personalize invitation images and send via WhatsApp Web.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    preview_parser = subparsers.add_parser("preview", help="Render one sample name onto the invitation.")
    preview_parser.add_argument("--name", required=True, help="Name to render, including Arabic if needed.")

    subparsers.add_parser("prepare", help="Generate personalized images for all Excel recipients.")
    subparsers.add_parser("send", help="Send personalized image and message through WhatsApp Web.")

    args = parser.parse_args()
    try:
        if args.command == "preview":
            command_preview(args.name)
        else:
            recipients = load_recipients()
            if args.command == "prepare":
                command_prepare(recipients)
            elif args.command == "send":
                command_send(recipients)
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()